"""Generate a presentation-ready architecture diagram and PowerPoint."""

import json
import os
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORT_DIR = os.path.join(BASE_DIR, "reports")
DELIVERABLE_DIR = os.path.join(BASE_DIR, "deliverables")
ARCH_PATH = os.path.join(DELIVERABLE_DIR, "architecture_big_data.png")
PPTX_PATH = os.path.join(DELIVERABLE_DIR, "presentation_semences.pptx")
TECH_REPORT_PATH = os.path.join(REPORT_DIR, "RAPPORT_TECHNIQUE.html")


def load_json(name):
    path = os.path.join(REPORT_DIR, name)
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as handle:
        return json.load(handle)


def architecture_diagram():
    fig, ax = plt.subplots(figsize=(16, 7), dpi=150)
    ax.set_xlim(0, 16)
    ax.set_ylim(0, 7)
    ax.axis("off")
    nodes = [
        (0.3, 4.4, 2.0, 1.2, "Sources\nCSV · météo · drones", "#2F6B4F"),
        (3.0, 4.4, 2.0, 1.2, "Ingestion\nNiFi · Sqoop · Kafka", "#2563A6"),
        (5.7, 4.4, 2.0, 1.2, "Raw\nHDFS / objet", "#74602B"),
        (8.4, 4.4, 2.0, 1.2, "Traitement\nSpark · ETL", "#8A3C3C"),
        (11.1, 4.4, 2.0, 1.2, "Trusted\nHive · HBase", "#5C4A86"),
        (13.8, 4.4, 1.9, 1.2, "Exposition\nR · ML · BI", "#1F6F78"),
        (5.7, 1.5, 2.6, 1.0, "Catalogue · Lineage\nAtlas cible / FTS local", "#444444"),
        (9.0, 1.5, 2.6, 1.0, "Qualité · Quarantaine\nContrats de données", "#444444"),
        (12.3, 1.5, 2.6, 1.0, "Monitoring · RBAC\nAPI · audit", "#444444"),
    ]
    for x, y, w, h, label, color in nodes:
        patch = FancyBboxPatch(
            (x, y), w, h, boxstyle="round,pad=0.04,rounding_size=0.08",
            facecolor=color, edgecolor="white", linewidth=1.5,
        )
        ax.add_patch(patch)
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center",
                color="white", fontsize=11, fontweight="bold")
    for x in [2.3, 5.0, 7.7, 10.4, 13.1]:
        ax.annotate("", xy=(x + 0.65, 5), xytext=(x, 5),
                    arrowprops={"arrowstyle": "->", "lw": 2, "color": "#333333"})
    for x in [7.0, 10.3, 13.6]:
        ax.annotate("", xy=(x, 4.35), xytext=(x, 2.55),
                    arrowprops={"arrowstyle": "-", "lw": 1.2, "color": "#777777"})
    ax.text(0.3, 6.35, "Architecture Big Data agricole - Prototype et cible",
            fontsize=20, fontweight="bold", color="#173B2C")
    ax.text(0.3, 6.0, "Flux de bout en bout avec R, gouvernance, sécurité et évolutivité",
            fontsize=11, color="#555555")
    fig.tight_layout()
    fig.savefig(ARCH_PATH, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(0.7))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(31, 84, 61)
    if subtitle:
        sub = slide.shapes.add_textbox(
            Inches(0.72), Inches(1.05), Inches(11.8), Inches(0.45)
        )
        p = sub.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(90, 90, 90)


def add_bullets(slide, bullets, top=1.6):
    box = slide.shapes.add_textbox(
        Inches(0.9), Inches(top), Inches(11.5), Inches(5.5)
    )
    frame = box.text_frame
    frame.word_wrap = True
    for index, text in enumerate(bullets):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = text
        p.font.size = Pt(20)
        p.space_after = Pt(14)


def presentation():
    prediction = load_json("prediction_report.json")
    quality = load_json("data_quality_report.json")
    kpi = load_json("kpi_globaux.json")
    metrics = prediction.get("classification_evaluation", {}).get("metrics", {})
    regression = prediction.get("metrics", {})
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Plateforme Big Data agricole SEMENCE",
              "Prototype fonctionnel pour le pôle Recherche & Développement")
    slide.shapes.add_picture(ARCH_PATH, Inches(0.8), Inches(1.55), width=Inches(11.7))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Contexte et objectifs")
    add_bullets(slide, [
        "Centraliser les données culturales et génétiques dans un Data Lake zoné.",
        "Garantir qualité, traçabilité, recherche et accès statistique depuis R.",
        "Produire indicateurs, graphes, géolocalisation et prédictions culturales.",
        "Préparer l'arrivée de données météo et d'images issues des drones.",
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Données et gouvernance")
    add_bullets(slide, [
        f"{kpi.get('nb_mesures', 0):,} mesures et {kpi.get('nb_essais', 0):,} essais.",
        f"{kpi.get('nb_materiels', 0):,} matériels génétiques catalogués.",
        "Catalogue de métadonnées, moteur de recherche FTS et lineage de bout en bout.",
        f"{quality.get('totals', {}).get('quarantined_rows', 0)} lignes mises en quarantaine.",
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Analyses et restitution")
    add_bullets(slide, [
        "Analyses Python, PySpark et Rscript avec ANOVA.",
        "Dashboard consolidé et carte interactive Folium.",
        "Requêtes analytiques Hive et Drill adaptées au prototype local.",
        "API protégée pour santé, catalogue, recherche et prédiction.",
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Modèle prédictif sans fuite")
    add_bullets(slide, [
        "Cible : rendement YD15QH; classe élevée à partir de 80 q/ha.",
        "GroupKFold par essai sur 2015-2017 et test final temporel 2018.",
        "Variables postérieures et identifiants mémorisables exclus.",
        f"MAE={regression.get('mae', 'N/A')} · RMSE={regression.get('rmse', 'N/A')} · R²={regression.get('r2', 'N/A')}.",
        f"Accuracy={metrics.get('accuracy', 'N/A')} · Précision={metrics.get('precision', 'N/A')} · Rappel={metrics.get('recall', 'N/A')} · F1={metrics.get('f1_score', 'N/A')}.",
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Exploitation et sécurité")
    add_bullets(slide, [
        "Checkpoint et reprise explicite avec python main.py --resume.",
        "RBAC local : chercheur, sélectionneur et opérateur.",
        "Journal d'audit API et historique des exécutions.",
        "Cible production : TLS, SSO, coffre de secrets, Prometheus/Grafana et haute disponibilité.",
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Limites et roadmap")
    add_bullets(slide, [
        "Acquérir météo, sol et conduite culturale pour améliorer la généralisation.",
        "Déployer les services distribués sur une infrastructure de préproduction.",
        "Valider les seuils et les résultats avec un expert agronome.",
        "Industrialiser sécurité, sauvegardes, supervision et tests de charge.",
    ])
    prs.save(PPTX_PATH)


def technical_report():
    if os.path.exists(TECH_REPORT_PATH):
        with open(TECH_REPORT_PATH, encoding="utf-8") as handle:
            if 'data-static-report="v2"' in handle.read(500):
                print("  Rapport technique éditorial v2 conservé")
                return
    prediction = load_json("prediction_report.json")
    quality = load_json("data_quality_report.json")
    hbase = load_json("hbase_summary.json")
    spark = load_json("spark_report.json")
    r_report = load_json("r_report.json")
    metrics = prediction.get("classification_evaluation", {}).get("metrics", {})
    hbase_load = hbase.get("tables", {}).get("results", {}).get("load", {})
    html = f"""<!doctype html><html lang="fr"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Rapport technique Semences</title>
<style>body{{font:16px Arial,sans-serif;max-width:1050px;margin:32px auto;padding:0 20px;color:#17231d}}
h1,h2{{color:#245d43}}table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #ccd5cf;padding:8px;text-align:left}}
.ok{{color:#176b3a;font-weight:bold}}.note{{background:#fff4d6;border-left:4px solid #c58b00;padding:12px}}</style></head><body>
<h1>Plateforme Big Data Semences - rapport technique vérifiable</h1>
<p>Généré le {datetime.now().isoformat(timespec='seconds')}. Ce livrable décrit un <strong>prototype fonctionnel local</strong>, pas un cluster distribué.</p>
<h2>Exécution réelle</h2><table><tr><th>Composant</th><th>Preuve</th><th>État</th></tr>
<tr><td>ETL pandas/Parquet</td><td>data/enriched/fact_table.parquet</td><td class="ok">Exécuté localement</td></tr>
<tr><td>Spark</td><td>reports/spark_report.json: {spark.get('engine','non vérifié')}</td><td class="ok">Moteur consigné</td></tr>
<tr><td>R</td><td>reports/r_report.json: {r_report.get('status','non vérifié')}</td><td class="ok">Rscript local</td></tr>
<tr><td>ML scikit-learn</td><td>Test temporel 2018, chevauchement groupes = {prediction.get('classification_evaluation',{}).get('train_test_group_overlap','N/A')}</td><td class="ok">Exécuté</td></tr>
</table>
<h2>Adaptations fonctionnelles locales</h2><p class="note">Sqoop, NiFi, HDFS, HBase, Kafka, Hive et Drill sont représentés par du code Python, des fichiers Parquet, SQLite, une file locale et DuckDB. Aucun service distribué correspondant n'est revendiqué comme installé.</p>
<h2>Contrôles chiffrés</h2><ul>
<li>Lignes contrôlées: {quality.get('totals',{}).get('rows','N/A')}; quarantaine: {quality.get('totals',{}).get('quarantined_rows','N/A')}.</li>
<li>HBase local: reçues {hbase_load.get('received','N/A')}, insérées {hbase_load.get('inserted','N/A')}, rejetées {hbase_load.get('rejected','N/A')}, collisions {hbase_load.get('collisions','N/A')}.</li>
<li>Classification 2018: accuracy {metrics.get('accuracy','N/A')}, précision {metrics.get('precision','N/A')}, rappel {metrics.get('recall','N/A')}, F1 {metrics.get('f1_score','N/A')}.</li></ul>
<h2>Limites et roadmap</h2><ul><li>Déploiement multi-nœuds, haute disponibilité, TLS/SSO et supervision centralisée restent à industrialiser.</li>
<li>Les données météo, sol et drones ne sont pas présentes; leurs connecteurs relèvent de la roadmap.</li>
<li>La date officielle de remise indiquée «XX» doit être confirmée.</li></ul>
</body></html>"""
    with open(TECH_REPORT_PATH, "w", encoding="utf-8") as handle:
        handle.write(html)


def run():
    print("=" * 80)
    print("  LIVRABLES - architecture et présentation")
    print("=" * 80)
    os.makedirs(DELIVERABLE_DIR, exist_ok=True)
    architecture_diagram()
    presentation()
    technical_report()
    report = {
        "status": "OK",
        "architecture": "deliverables/architecture_big_data.png",
        "presentation": "deliverables/presentation_semences.pptx",
        "technical_report": "reports/RAPPORT_TECHNIQUE.html",
        "slides": 7,
        "timestamp": datetime.now().isoformat(),
    }
    with open(os.path.join(REPORT_DIR, "deliverables_report.json"), "w",
              encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2)
    print(f"  Diagramme: {ARCH_PATH}")
    print(f"  Présentation: {PPTX_PATH}")
    return report


if __name__ == "__main__":
    run()
